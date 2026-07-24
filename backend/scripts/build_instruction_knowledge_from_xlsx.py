#!/usr/bin/env python3
"""
说明书知识库构建脚本：把 Excel 清洗结果转换为结构化 JSON 知识库。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from hashlib import sha1
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - exercised by runtime environment
    raise SystemExit(
        "Missing dependency: openpyxl. Install backend requirements first: "
        "pip install -r backend/requirements.txt"
    ) from exc


BACKEND_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_KNOWLEDGE_OUTPUT = BACKEND_ROOT / "data" / "processed" / "instruction_knowledge_from_xlsx.json"
DEFAULT_KG_OUTPUT = BACKEND_ROOT / "data" / "processed" / "instruction_kg_candidates.jsonl"
DEFAULT_REPORT_OUTPUT = BACKEND_ROOT / "data" / "processed" / "instruction_build_report.json"


SECTION_COLUMNS = {
    "适应症": "适应症",
    "不良反应": "不良反应",
    "用法用量": "用法用量",
    "禁忌": "禁忌",
    "注意事项": "警告与注意事项",
    "孕妇及哺乳期妇女用药": "孕妇及哺乳期妇女用药",
    "儿童用药": "儿童用药",
    "老人用药": "老年用药",
    "药物相互作用": "药物相互作用",
    "药理毒理": "药理毒理",
    "药代动力学": "药代动力学",
}

SAFETY_SECTION_TITLES = {
    "禁忌",
    "警告与注意事项",
    "孕妇及哺乳期妇女用药",
    "儿童用药",
    "老年用药",
    "药物相互作用",
}

HIGH_TERMS = ("禁用", "禁忌", "严禁", "禁止", "不得", "避免合用", "避免使用", "严禁并用")
MEDIUM_TERMS = (
    "慎用",
    "需注意",
    "应注意",
    "监测",
    "调整剂量",
    "减量",
    "风险增加",
    "风险增高",
    "不宜",
    "医师指导",
    "医生指导",
    "药师指导",
    "专业医师指导",
    "遵医嘱",
    "成人监护",
)
NON_ACTIONABLE_TERMS = (
    "无特别注意事项",
    "无特殊要求",
    "无特殊注意",
    "尚不明确",
    "尚未明确",
    "尚无资料",
    "资料不详",
    "未进行该项实验",
    "未进行相关研究",
)

POPULATION_PATTERNS = {
    "儿童": ("儿童", "小儿", "婴幼儿", "新生儿", "18岁以下", "18 岁以下"),
    "老年人": ("老人", "老年", "高龄"),
    "妊娠/备孕": ("孕妇", "妊娠", "孕期", "备孕", "胎儿"),
    "哺乳期": ("哺乳", "乳汁", "乳母"),
}

CONDITION_PATTERNS = {
    "肝功能不全": ("肝功能不全", "肝功能障碍", "严重肝", "活动性肝病"),
    "肾功能不全": ("肾功能不全", "肾功能障碍", "严重肾", "透析"),
    "消化道出血/溃疡": ("消化道出血", "胃肠道出血", "消化道溃疡", "胃溃疡", "活动性出血"),
    "过敏": ("过敏", "过敏史"),
    "心力衰竭": ("心力衰竭", "心衰"),
    "哮喘": ("哮喘",),
    "糖尿病": ("糖尿病",),
}

DRUG_CLASS_PATTERNS = {
    "抗凝药": ("抗凝药", "华法林", "肝素"),
    "抗血小板药": ("抗血小板", "阿司匹林", "氯吡格雷"),
    "NSAIDs": ("NSAIDs", "非甾体抗炎药", "解热镇痛药"),
    "CYP3A4抑制剂": ("CYP3A4抑制剂", "CYP3A4 抑制剂"),
    "CYP2D6抑制剂": ("CYP2D6抑制剂", "CYP2D6 抑制剂"),
    "CYP2C19抑制剂": ("CYP2C19抑制剂", "CYP2C19 抑制剂"),
    "MAOIs": ("MAOIs", "单胺氧化酶抑制剂"),
    "中枢抑制药": ("中枢抑制药", "镇静催眠药", "酒精"),
}

RISK_EVENT_PATTERNS = {
    "出血": ("出血", "INR"),
    "低血压": ("低血压",),
    "高钾血症": ("高钾", "血钾"),
    "肝损伤": ("肝损害", "肝损伤", "转氨酶"),
    "肾损伤": ("肾损害", "肾损伤", "肾功能恶化"),
    "横纹肌溶解": ("横纹肌溶解", "肌病", "肌痛"),
    "骨髓抑制": ("骨髓抑制", "白细胞减少", "血小板减少"),
    "过敏反应": ("过敏反应", "休克", "荨麻疹"),
}

SHORT_STRUCTURAL_SECTIONS = {
    "禁忌",
    "孕妇及哺乳期妇女用药",
    "儿童用药",
    "老年用药",
}
LONG_STRUCTURAL_SECTIONS = {
    "警告与注意事项",
    "药物相互作用",
    "不良反应",
    "用法用量",
    "药理毒理",
    "药代动力学",
}
STRUCTURAL_CHUNK_MAX_CHARS = 360
STRUCTURAL_CHUNK_OVERLAP_SENTENCES = 1


@dataclass(frozen=True)
class SourceRef:
    row: int
    approval_no: str
    source_url: str


@dataclass
class SectionAggregate:
    title: str
    texts: list[str]
    source_rows: list[int]
    text_hashes: set[str]


@dataclass
class KGCandidate:
    subject: str
    relation: str
    object: str
    object_type: str
    risk_level: str
    evidence_text: str
    section: str
    source_rows: list[int]
    source_urls: list[str]
    confidence: float


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    text = text.replace("\u3000", " ")
    text = text.replace("|", "。")
    text = re.sub(r"[?？]{2,}", "？", text)
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"([。；;])\s*", r"\1", text)
    return text.strip(" ;；。")


def normalize_name(value: Any) -> str:
    text = normalize_text(value)
    text = re.sub(r"\s+", "", text)
    return text


def split_sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[。！？；;])|(?:\s*\d+[.)、]\s*)", text)
    return [part.strip(" ;；。") for part in parts if len(part.strip(" ;；。")) >= 4]


def split_numbered_items(text: str) -> list[str]:
    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    parts = re.split(r"(?:^|[。；;]\s*|\n+)\s*(?:\d+|[一二三四五六七八九十]+)[.、)]\s*", normalized)
    items = [part.strip(" ;；。") for part in parts if len(part.strip(" ;；。")) >= 4]
    if len(items) > 1:
        return items
    return split_sentences(normalized) or [normalized]


def pack_text_units(units: list[str], max_chars: int = STRUCTURAL_CHUNK_MAX_CHARS) -> list[str]:
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for unit in units:
        if len(unit) > max_chars:
            if current:
                chunks.append("。".join(current).strip("。") + "。")
                current = []
                current_len = 0
            start = 0
            while start < len(unit):
                chunks.append(unit[start : start + max_chars].strip())
                start += max_chars
            continue

        extra_len = len(unit) + (1 if current else 0)
        if current and current_len + extra_len > max_chars:
            chunks.append("。".join(current).strip("。") + "。")
            current = current[-STRUCTURAL_CHUNK_OVERLAP_SENTENCES:] if STRUCTURAL_CHUNK_OVERLAP_SENTENCES else []
            current_len = sum(len(item) for item in current)

        current.append(unit)
        current_len += extra_len

    if current:
        chunks.append("。".join(current).strip("。") + "。")
    return [chunk for chunk in chunks if len(chunk.strip()) >= 4]


def risk_terms_from_text(text: str) -> list[str]:
    terms = [term for term in (*HIGH_TERMS, *MEDIUM_TERMS) if term in text]
    return list(dict.fromkeys(terms))


def make_section_chunks(section_title: str, texts: list[str], source_rows: list[int]) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for variant_index, text in enumerate(texts, start=1):
        if section_title in SHORT_STRUCTURAL_SECTIONS or len(text) <= STRUCTURAL_CHUNK_MAX_CHARS:
            chunk_texts = [text]
            chunk_type = "section"
        else:
            units = split_numbered_items(text)
            chunk_texts = pack_text_units(units)
            chunk_type = "clause" if section_title in LONG_STRUCTURAL_SECTIONS else "section_part"

        for chunk_index, chunk_text in enumerate(chunk_texts):
            chunks.append(
                {
                    "text": chunk_text,
                    "chunk_type": chunk_type,
                    "variant_index": variant_index,
                    "chunk_index": chunk_index,
                    "source_rows": source_rows,
                    "population_tags": match_named_objects(chunk_text, POPULATION_PATTERNS),
                    "condition_tags": match_named_objects(chunk_text, CONDITION_PATTERNS),
                    "risk_terms": risk_terms_from_text(chunk_text),
                }
            )
    return chunks


def risk_level_from_text(text: str) -> str:
    if any(term in text for term in HIGH_TERMS):
        return "High"
    if any(term in text for term in MEDIUM_TERMS):
        return "Medium"
    return "Low"


def is_non_actionable_statement(text: str) -> bool:
    if any(term in text for term in HIGH_TERMS):
        return False
    if any(term in text for term in ("慎用", "风险增加", "风险增高", "调整剂量", "减量", "医师指导", "医生指导", "药师指导", "遵医嘱")):
        return False
    return any(term in text for term in NON_ACTIONABLE_TERMS)


def relation_for_section(section: str, text: str) -> str | None:
    if is_non_actionable_statement(text):
        return None
    if section == "药物相互作用":
        return "INTERACTS_WITH"
    if any(term in text for term in HIGH_TERMS):
        return "CONTRAINDICATED_FOR"
    if any(term in text for term in MEDIUM_TERMS):
        return "USE_WITH_CAUTION_IN"
    return None


def text_hash(text: str) -> str:
    return sha1(text.encode("utf-8")).hexdigest()


def add_unique_text(bucket: SectionAggregate, text: str, row: int, min_chars: int) -> None:
    if len(text) < min_chars:
        return
    digest = text_hash(text)
    if digest not in bucket.text_hashes:
        bucket.texts.append(text)
        bucket.text_hashes.add(digest)
    if row not in bucket.source_rows:
        bucket.source_rows.append(row)


def iter_rows(path: Path, sheet_name: str | None = None, max_rows: int | None = None) -> Iterable[tuple[int, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(item).strip() if item is not None else "" for item in header_row]
    for excel_row, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if max_rows is not None and excel_row > max_rows + 1:
            break
        yield excel_row, dict(zip(headers, values))


def build_aliases(rows: list[dict[str, Any]]) -> list[str]:
    aliases: set[str] = set()
    for row in rows:
        generic = normalize_name(row.get("通用名称"))
        brand_name = normalize_name(row.get("商品名称"))
        if brand_name and brand_name != generic:
            aliases.add(brand_name)

        title = normalize_name(row.get("标题"))
        title_alias = re.sub(r"(说明书|药品说明书)$", "", title)
        if title_alias and title_alias != generic and generic and generic in title_alias:
            aliases.add(title_alias)
    return sorted(aliases)


def collect_known_drug_names(rows_by_drug: dict[str, list[dict[str, Any]]]) -> dict[str, str]:
    names: dict[str, str] = {}
    for drug, rows in rows_by_drug.items():
        names[drug] = drug
        for alias in build_aliases(rows):
            if len(alias) >= 2:
                names[alias] = drug
    return names


def index_known_drug_names(known_names: dict[str, str]) -> dict[str, list[tuple[str, str]]]:
    indexed: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for name, canonical in known_names.items():
        if len(name) < 2:
            continue
        indexed[name[0]].append((name, canonical))
    for names in indexed.values():
        names.sort(key=lambda item: len(item[0]), reverse=True)
    return indexed


def match_named_objects(text: str, patterns: dict[str, tuple[str, ...]]) -> list[str]:
    return [name for name, terms in patterns.items() if any(term in text for term in terms)]


def extract_interaction_objects(
    text: str,
    subject: str,
    known_name_index: dict[str, list[tuple[str, str]]],
) -> list[tuple[str, str]]:
    objects: dict[str, str] = {}
    candidate_names: list[tuple[str, str]] = []
    for char in set(text):
        candidate_names.extend(known_name_index.get(char, []))
    for name, canonical in candidate_names:
        if canonical == subject or len(name) < 2:
            continue
        if name in text:
            objects[canonical] = "Drug"
    for class_name in match_named_objects(text, DRUG_CLASS_PATTERNS):
        objects[class_name] = "DrugClass"
    return sorted(objects.items())


def make_candidate(
    *,
    subject: str,
    relation: str,
    object_name: str,
    object_type: str,
    evidence_text: str,
    section: str,
    refs: list[SourceRef],
    confidence: float,
) -> KGCandidate:
    return KGCandidate(
        subject=subject,
        relation=relation,
        object=object_name,
        object_type=object_type,
        risk_level=risk_level_from_text(evidence_text),
        evidence_text=evidence_text[:600],
        section=section,
        source_rows=sorted({ref.row for ref in refs}),
        source_urls=sorted({ref.source_url for ref in refs if ref.source_url}),
        confidence=confidence,
    )


def extract_kg_candidates(
    drug: str,
    sections: list[SectionAggregate],
    refs_by_row: dict[int, SourceRef],
    known_name_index: dict[str, list[tuple[str, str]]],
) -> list[KGCandidate]:
    candidates: list[KGCandidate] = []
    seen: set[tuple[str, str, str, str, str]] = set()

    for section in sections:
        if section.title not in SAFETY_SECTION_TITLES:
            continue
        refs = [refs_by_row[row] for row in section.source_rows if row in refs_by_row]
        for text in section.texts:
            for sentence in split_sentences(text):
                relation = relation_for_section(section.title, sentence)
                if not relation:
                    continue

                if section.title == "药物相互作用":
                    for object_name, object_type in extract_interaction_objects(sentence, drug, known_name_index):
                        key = (drug, relation, object_name, section.title, sentence[:80])
                        if key in seen:
                            continue
                        seen.add(key)
                        candidates.append(
                            make_candidate(
                                subject=drug,
                                relation=relation,
                                object_name=object_name,
                                object_type=object_type,
                                evidence_text=sentence,
                                section=section.title,
                                refs=refs,
                                confidence=0.72 if object_type == "Drug" else 0.62,
                            )
                        )
                    continue

                targets: list[tuple[str, str]] = []
                targets.extend((name, "Population") for name in match_named_objects(sentence, POPULATION_PATTERNS))
                targets.extend((name, "Condition") for name in match_named_objects(sentence, CONDITION_PATTERNS))
                targets.extend((name, "Risk") for name in match_named_objects(sentence, RISK_EVENT_PATTERNS))

                for object_name, object_type in targets:
                    key = (drug, relation, object_name, section.title, sentence[:80])
                    if key in seen:
                        continue
                    seen.add(key)
                    candidates.append(
                        make_candidate(
                            subject=drug,
                            relation=relation,
                            object_name=object_name,
                            object_type=object_type,
                            evidence_text=sentence,
                            section=section.title,
                            refs=refs,
                            confidence=0.68,
                        )
                    )
    return candidates


def resolve_xlsx_inputs(inputs: Iterable[Path]) -> list[Path]:
    paths: list[Path] = []
    for item in inputs:
        if item.is_dir():
            paths.extend(
                sorted(
                    path
                    for path in item.rglob("*.xlsx")
                    if not path.name.startswith("~$")
                )
            )
            continue
        paths.append(item)

    unique: list[Path] = []
    seen: set[Path] = set()
    for path in paths:
        resolved = path.expanduser().resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        unique.append(resolved)
    if not unique:
        raise ValueError("No .xlsx input files found.")
    return unique


def iter_input_rows(
    paths: Iterable[Path],
    sheet_name: str | None = None,
    max_rows: int | None = None,
) -> Iterable[tuple[int, dict[str, Any]]]:
    total = 0
    for path in paths:
        for excel_row, row in iter_rows(path, sheet_name=sheet_name, max_rows=None):
            if max_rows is not None and total >= max_rows:
                return
            total += 1
            row["__source_file"] = str(path)
            row["__excel_row"] = excel_row
            yield total, row


def build_records_from_paths(
    paths: Iterable[Path],
    sheet_name: str | None,
    max_rows: int | None,
    min_section_chars: int,
) -> tuple[list[dict[str, Any]], list[KGCandidate], dict[str, Any]]:
    input_paths = resolve_xlsx_inputs(paths)
    rows_by_drug: dict[str, list[dict[str, Any]]] = defaultdict(list)
    refs_by_drug: dict[str, dict[int, SourceRef]] = defaultdict(dict)
    skipped = Counter()
    rows_by_file = Counter()

    for row_num, row in iter_input_rows(input_paths, sheet_name=sheet_name, max_rows=max_rows):
        source_file = normalize_text(row.get("__source_file"))
        if source_file:
            rows_by_file[source_file] += 1
        drug = normalize_name(row.get("通用名称")) or normalize_name(row.get("标题"))
        if not drug:
            skipped["missing_drug_name"] += 1
            continue
        row["__source_row"] = row_num
        rows_by_drug[drug].append(row)
        refs_by_drug[drug][row_num] = SourceRef(
            row=row_num,
            approval_no=normalize_text(row.get("批准文号") or row.get("编号")),
            source_url=normalize_text(row.get("标题链接")),
        )

    known_names = collect_known_drug_names(rows_by_drug)
    known_name_index = index_known_drug_names(known_names)
    records: list[dict[str, Any]] = []
    kg_candidates: list[KGCandidate] = []
    section_counts = Counter()

    for drug in sorted(rows_by_drug):
        rows = rows_by_drug[drug]
        sections_by_title: dict[str, SectionAggregate] = {
            title: SectionAggregate(title=title, texts=[], source_rows=[], text_hashes=set())
            for title in dict.fromkeys(SECTION_COLUMNS.values())
        }
        for row in rows:
            source_row = int(row.get("__source_row") or 0)
            for column, section_title in SECTION_COLUMNS.items():
                text = normalize_text(row.get(column))
                add_unique_text(sections_by_title[section_title], text, source_row, min_section_chars)

        sections: list[dict[str, Any]] = []
        section_aggregates: list[SectionAggregate] = []
        for title, aggregate in sections_by_title.items():
            if not aggregate.texts:
                continue
            content = "\n".join(f"{idx + 1}. {text}" for idx, text in enumerate(aggregate.texts))
            sections.append(
                {
                    "title": title,
                    "content": content,
                    "source_rows": aggregate.source_rows,
                    "variant_count": len(aggregate.texts),
                    "chunks": make_section_chunks(title, aggregate.texts, aggregate.source_rows),
                }
            )
            section_aggregates.append(aggregate)
            section_counts[title] += 1

        if not sections:
            skipped["empty_sections"] += 1
            continue

        aliases = build_aliases(rows)
        source_urls = sorted(
            {
                normalize_text(row.get("标题链接"))
                for row in rows
                if normalize_text(row.get("标题链接"))
            }
        )
        approval_numbers = sorted(
            {
                normalize_text(row.get("批准文号") or row.get("编号"))
                for row in rows
                if normalize_text(row.get("批准文号") or row.get("编号"))
            }
        )
        source_files = sorted(
            {
                normalize_text(row.get("__source_file"))
                for row in rows
                if normalize_text(row.get("__source_file"))
            }
        )
        records.append(
            {
                "drug": drug,
                "aliases": aliases,
                "source": "药品说明书数据库 XLSX 聚合",
                "source_urls": source_urls[:50],
                "source_files": source_files,
                "approval_numbers": approval_numbers[:100],
                "source_row_count": len(rows),
                "sections": sections,
                "interactions": [],
            }
        )
        kg_candidates.extend(
            extract_kg_candidates(
                drug=drug,
                sections=section_aggregates,
                refs_by_row=refs_by_drug[drug],
                known_name_index=known_name_index,
            )
        )

    report = {
        "inputs": [str(path) for path in input_paths],
        "input_file_count": len(input_paths),
        "rows_by_file": dict(rows_by_file),
        "rows_read": sum(len(rows) for rows in rows_by_drug.values()),
        "drug_count": len(records),
        "kg_candidate_count": len(kg_candidates),
        "section_counts": dict(section_counts),
        "chunk_count": sum(
            len(section.get("chunks", []))
            for record in records
            for section in record.get("sections", [])
        ),
        "chunking": {
            "strategy": "section_first_then_clause_chunks",
            "max_chunk_chars": STRUCTURAL_CHUNK_MAX_CHARS,
            "short_structural_sections": sorted(SHORT_STRUCTURAL_SECTIONS),
            "long_structural_sections": sorted(LONG_STRUCTURAL_SECTIONS),
        },
        "skipped": dict(skipped),
    }
    return records, kg_candidates, report


def build_records(path: Path, sheet_name: str | None, max_rows: int | None, min_section_chars: int) -> tuple[list[dict[str, Any]], list[KGCandidate], dict[str, Any]]:
    return build_records_from_paths(
        paths=[path],
        sheet_name=sheet_name,
        max_rows=max_rows,
        min_section_chars=min_section_chars,
    )


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_jsonl(path: Path, rows: Iterable[Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        for row in rows:
            file.write(json.dumps(asdict(row), ensure_ascii=False) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build SafeMeds RAG records and KG candidates from Chinese drug instruction XLSX files."
    )
    parser.add_argument("input", type=Path, nargs="+", help="Input .xlsx file(s), or directories containing .xlsx files.")
    parser.add_argument("--sheet", default=None, help="Sheet name. Defaults to the first sheet.")
    parser.add_argument("--knowledge-output", type=Path, default=DEFAULT_KNOWLEDGE_OUTPUT)
    parser.add_argument("--kg-output", type=Path, default=DEFAULT_KG_OUTPUT)
    parser.add_argument("--report-output", type=Path, default=DEFAULT_REPORT_OUTPUT)
    parser.add_argument("--max-rows", type=int, default=None, help="Optional row limit for dry runs.")
    parser.add_argument("--min-section-chars", type=int, default=4)
    args = parser.parse_args()

    records, kg_candidates, report = build_records_from_paths(
        paths=args.input,
        sheet_name=args.sheet,
        max_rows=args.max_rows,
        min_section_chars=args.min_section_chars,
    )
    write_json(args.knowledge_output, records)
    write_jsonl(args.kg_output, kg_candidates)
    write_json(args.report_output, report)

    print(
        "Instruction XLSX build completed: "
        f"rows={report['rows_read']} drugs={report['drug_count']} "
        f"kg_candidates={report['kg_candidate_count']}"
    )
    print(f"knowledge: {args.knowledge_output}")
    print(f"kg_candidates: {args.kg_output}")
    print(f"report: {args.report_output}")


if __name__ == "__main__":
    main()
